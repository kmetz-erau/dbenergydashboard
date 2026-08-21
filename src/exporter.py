from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .analytics import AnalyticsBundle

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _safe_sheet_name(name: str) -> str:
    return name[:31]


def _format_workbook(buffer: BytesIO) -> BytesIO:
    buffer.seek(0)
    wb = load_workbook(buffer)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        for col_idx, column_cells in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), start=1):
            max_len = 0
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 32)

        header_index = {cell.value: cell.column for cell in ws[1]}
        for name, idx in header_index.items():
            if not isinstance(name, str):
                continue
            lname = name.lower()
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row, idx)
                if "pct" in lname or "percent" in lname:
                    cell.number_format = "0.0%"
                elif "cost" in lname or "amount" in lname or "charge" in lname:
                    cell.number_format = '$#,##0.00;[Red]($#,##0.00);-'
                elif "eui" in lname:
                    cell.number_format = "0.0"
                elif "demand" in lname and "charge" not in lname:
                    cell.number_format = "#,##0.0"
                elif "intensity" in lname:
                    cell.number_format = "0.00"
                elif "kwh" in lname or "therm" in lname:
                    cell.number_format = "#,##0.0"
                elif "sqft" in lname or lname.endswith("sf"):
                    cell.number_format = "#,##0"
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def build_analysis_workbook(
    bundle: AnalyticsBundle,
    buildings: pd.DataFrame,
    meters: pd.DataFrame,
    exceptions: pd.DataFrame | None = None,
    gas_bundle=None,
    gas_accounts: pd.DataFrame | None = None,
    combined_snapshot: pd.DataFrame | None = None,
) -> bytes:
    """Build one analysis workbook.

    Existing FPL sheets remain unchanged. Gas and combined-site sheets are appended
    only when gas data/reference data are supplied.
    """
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        # Existing FPL outputs.
        bundle.bills.to_excel(writer, sheet_name="Processed Bill History", index=False)
        bundle.allocated_bills.to_excel(writer, sheet_name="Allocated Bills", index=False)
        bundle.building_monthly.to_excel(writer, sheet_name="Building Monthly", index=False)
        bundle.building_snapshot.to_excel(writer, sheet_name="Building Snapshot", index=False)
        bundle.campus_monthly.to_excel(writer, sheet_name="Campus Monthly", index=False)
        bundle.demand_details.to_excel(writer, sheet_name="Demand History", index=False)
        bundle.allocated_demand.to_excel(writer, sheet_name="Allocated Demand", index=False)
        bundle.building_demand_monthly.to_excel(writer, sheet_name="Building Demand", index=False)
        bundle.building_demand_snapshot.to_excel(writer, sheet_name="Demand Snapshot", index=False)
        bundle.unmapped_demand_meters.to_excel(writer, sheet_name="Unmapped Demand", index=False)
        bundle.unmapped_meters.to_excel(writer, sheet_name="Unmapped Meters", index=False)
        bundle.meter_suggestions.to_excel(writer, sheet_name="Meter Suggestions", index=False)
        bundle.shared_meter_issues.to_excel(writer, sheet_name="Shared Meter Issues", index=False)
        bundle.missing_area.to_excel(writer, sheet_name="Missing Area", index=False)
        if exceptions is not None:
            exceptions.to_excel(writer, sheet_name="Exceptions", index=False)
        buildings.to_excel(writer, sheet_name="Buildings Reference", index=False)
        meters.to_excel(writer, sheet_name="Meters Reference", index=False)

        # Additive gas outputs.
        if gas_bundle is not None:
            gas_bundle.history.to_excel(writer, sheet_name="Gas History", index=False)
            gas_bundle.enriched_history.to_excel(writer, sheet_name="Gas Enriched", index=False)
            gas_bundle.usage_group_monthly.to_excel(writer, sheet_name="Gas Usage Groups", index=False)
            gas_bundle.building_monthly.to_excel(writer, sheet_name="Gas Building Monthly", index=False)
            gas_bundle.building_snapshot.to_excel(writer, sheet_name="Gas Snapshot", index=False)
            gas_bundle.campus_monthly.to_excel(writer, sheet_name="Gas Campus Monthly", index=False)
            gas_bundle.provider_summary.to_excel(writer, sheet_name="Gas Providers", index=False)
            gas_bundle.unmapped_accounts.to_excel(writer, sheet_name="Gas Unmapped", index=False)
            gas_bundle.fixed_charge_rows.to_excel(writer, sheet_name="Gas Fixed Charges", index=False)
        if gas_accounts is not None:
            gas_accounts.to_excel(writer, sheet_name="Gas Accounts Ref", index=False)
        if combined_snapshot is not None:
            combined_snapshot.to_excel(writer, sheet_name="Total Site EUI", index=False)
    return _format_workbook(buffer).getvalue()
